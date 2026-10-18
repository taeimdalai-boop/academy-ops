import email
import email.policy
import os
import re
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.payslip.company import load_company
from app.payslip.ledger import load_ledger
from app.payslip.render import render_html

S = "samples/payslip/workflowlab"
LEDGER = f"{S}/급여대장_2026-09.xlsx"


@pytest.fixture(scope="module")
def company():
    return load_company(f"{S}/company.yaml")


def test_ledger_rows_and_totals(company):
    slips, issues = load_ledger(LEDGER, company)
    assert [s.employee_no for s in slips] == ["WL-001", "WL-002", "WL-003", "WL-004", "WL-005"]
    s4 = slips[3]
    assert s4.gross == Decimal("1770000")
    assert s4.deductions == Decimal("83040")          # 두루누리 음수 포함
    assert s4.net == Decimal("1686960")
    assert issues == []


def test_items_keep_config_order_and_pad_to_ten(company):
    slips, _ = load_ledger(LEDGER, company)
    pay = slips[0].pay_rows()
    assert [r[0] for r in pay[:3]] == ["기본급", "식대", "자가운전보조금"]
    assert len(pay) == 10 and pay[-1] == ("-", None)


def test_total_mismatch_blocks_employee(company, tmp_path):
    wb = load_workbook(LEDGER)
    ws = wb["급여대장"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("실지급액") + 1, value=1)   # WL-001 실지급액 틀리게
    bad = tmp_path / "bad.xlsx"
    wb.save(bad)
    slips, issues = load_ledger(str(bad), company)
    assert "WL-001" not in [s.employee_no for s in slips]
    assert any("WL-001" in i and "실지급액" in i for i in issues)


def test_missing_item_column_stops(company, tmp_path):
    wb = load_workbook(LEDGER)
    ws = wb["급여대장"]
    header = [c.value for c in ws[1]]
    ws.delete_cols(header.index("식대") + 1)
    bad = tmp_path / "bad.xlsx"
    wb.save(bad)
    from app.xlsx import InputError
    with pytest.raises(InputError, match="식대"):
        load_ledger(str(bad), company)


def test_html_has_brand_and_values(company):
    slips, _ = load_ledger(LEDGER, company)
    html = render_html(slips[3], company, base_dir=os.path.abspath(S))
    for text in ["월 급 명 세 서", "PAY STATEMENT", "워크플로우랩", "정하은", "2026년 09월", "2026년 10월 10일",
                 "-54,720", "1,770,000", "83,040", "1,686,960", "두루누리", "함께 성장해 주셔서", "김태희 소장"]:
        assert text in html, text
    assert "우리컴퓨터" not in html


def test_cli_payslip_end_to_end(tmp_path):
    from app.payslip.cli import run_payslips
    code = run_payslips(LEDGER, f"{S}/company.yaml", str(tmp_path))
    assert code == 0
    pdfs = sorted(os.listdir(tmp_path / "pdf"))
    assert pdfs == sorted(f"{n}_급여명세서_2026-09.pdf" for n in ["박지훈", "이수아", "최도윤", "정하은", "한예린"])
    import fitz
    doc = fitz.open(tmp_path / "pdf" / "정하은_급여명세서_2026-09.pdf")
    assert len(doc) == 1
    text = doc[0].get_text()
    assert "1,686,960" in text and "워크플로우랩" in text
    msg = email.message_from_bytes((tmp_path / "mail" / "WL-004.eml").read_bytes(), policy=email.policy.default)
    parts = [p for p in msg.walk() if p.get_filename()]
    assert parts and parts[0].get_filename().endswith(".pdf")
    body = next(p for p in msg.walk() if p.get_content_type() == "text/plain").get_content()
    assert not re.search(r"\d{1,3}(,\d{3})+|\d+\s*원", body)
