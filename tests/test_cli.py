import json
import os
import re

from openpyxl import load_workbook

from app.cli import run


def test_cli_end_to_end(tmp_path):
    code = run("2026-09", "samples/2026-09", str(tmp_path))
    assert code == 2                                   # T04 최저시급 미달

    wb = load_workbook(tmp_path / "급여계산표_2026-09.xlsx")
    assert wb.sheetnames == ["요약", "명세항목", "근태집계", "확인필요"]

    # 계산근거 = 금액이 어떻게 나왔는지 검산하는 내부 자료(직원ID로 식별)
    basis = sorted(os.listdir(tmp_path / "계산근거"))
    assert basis == ["A01_2026-09.pdf", "A02_2026-09.pdf", "T01_2026-09.pdf", "T02_2026-09.pdf", "T03_2026-09.pdf"]
    assert all((tmp_path / "계산근거" / p).stat().st_size > 2000 for p in basis)

    # 명세서 = 직원에게 주는 문서(실물 양식, 성명으로 식별). company.yaml이 있을 때만 생성된다.
    slips = sorted(f for f in os.listdir(tmp_path / "명세서") if f.endswith(".pdf"))
    assert slips == ["강사A_급여명세서_2026-09.pdf", "강사B_급여명세서_2026-09.pdf", "강사C_급여명세서_2026-09.pdf",
                     "행정A_급여명세서_2026-09.pdf", "행정B_급여명세서_2026-09.pdf"]
    assert all((tmp_path / "명세서" / p).stat().st_size > 2000 for p in slips)
    # 계산 중단된 T04(강사D)는 명세서를 만들지 않는다
    assert not any("강사D" in f for f in slips)

    for f in os.listdir(tmp_path / "메일초안"):
        body = (tmp_path / "메일초안" / f).read_text(encoding="utf-8")
        assert not re.search(r"\d{1,3}(,\d{3})+|\d+\s*원", body), f

    md = (tmp_path / "확인필요_2026-09.md").read_text(encoding="utf-8")
    for word in ["S999", "9999", "T04", "대체", "8시간", "지각", "4대보험"]:
        assert word in md, word

    meta = json.loads((tmp_path / "run.json").read_text(encoding="utf-8"))
    assert meta["status"] == "CALCULATED"
    assert meta["totals"] == {"gross": 6971000, "deductions": 100000, "net": 6871000}
    assert meta["blocked"] == ["T04"]


def test_missing_input_file(tmp_path):
    code = run("2026-09", str(tmp_path / "없는폴더"), str(tmp_path / "out"))
    assert code == 1
