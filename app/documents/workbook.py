"""급여 계산표(엑셀): 요약 · 명세항목 · 근태집계 · 확인필요."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.payroll.engine import describe

HEAD = PatternFill("solid", fgColor="1A3A5C")
WON = '#,##0'


def _sheet(wb, title, header, rows, widths, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = WON
    ws.freeze_panes = "A2"
    return ws


def write_workbook(path, statements, attendance, issues, blocked):
    wb = Workbook()
    wb.remove(wb.active)
    summary = [[s.employee_id, s.name, int(s.gross), int(s.deductions), int(s.net), "계산 완료"] for s in statements.values()]
    summary += [[emp, "", None, None, None, "계산 중단 (확인필요 참고)"] for emp in blocked]
    summary.append(["합계", "", sum(int(s.gross) for s in statements.values()),
                    sum(int(s.deductions) for s in statements.values()),
                    sum(int(s.net) for s in statements.values()), ""])
    _sheet(wb, "요약", ["직원ID", "이름", "지급액", "공제액", "실지급액", "상태"], summary, [10, 10, 14, 14, 14, 26], (3, 4, 5))

    lines = []
    for s in statements.values():
        for l in s.lines:
            lines.append([s.employee_id, s.name, l.item, l.side, float(l.quantity), l.unit,
                          float(l.unit_rate) if l.unit_rate is not None else None, int(l.amount), describe(l), l.basis])
    _sheet(wb, "명세항목", ["직원ID", "이름", "항목", "구분", "수량", "단위", "단가", "금액", "계산", "근거"], lines,
           [8, 8, 14, 6, 8, 9, 10, 12, 34, 44], (7, 8))

    att_rows = []
    for emp, m in sorted(attendance.items()):
        for d in m.days:
            att_rows.append([emp, d.day, d.check_in.strftime("%H:%M") if d.check_in else "",
                             d.check_out.strftime("%H:%M") if d.check_out else "", d.break_minutes,
                             round(d.work_minutes / 60, 2), d.note])
    _sheet(wb, "근태집계", ["직원ID", "날짜", "출근", "퇴근", "외출(분)", "근무(시간)", "비고"], att_rows, [8, 12, 8, 8, 9, 10, 36])
    for row in wb["근태집계"].iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd"

    iss = [[i.category, i.employee_id or "", i.day or "", i.message] for i in issues]
    _sheet(wb, "확인필요", ["분류", "직원ID", "날짜", "내용"], iss, [12, 8, 12, 90])
    wb.save(path)
